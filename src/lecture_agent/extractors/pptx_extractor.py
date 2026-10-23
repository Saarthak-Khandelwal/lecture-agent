from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Iterator

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

_DIAGRAM_URI = "drawingml/2006/diagram"


def _run_text(run) -> str:
    text = run.text
    try:
        addr = run.hyperlink.address
    except Exception:
        addr = None
    if addr and text.strip():
        text += f" ({addr})"  # keeps URLs from the slides available for the references step
    return text


def _shape_text(shape) -> list[str]:
    lines: list[str] = []
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for sub in shape.shapes:
            lines.extend(_shape_text(sub))
        return lines
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            cells = [c.text.strip().replace("\n", " ") for c in row.cells]
            lines.append("| " + " | ".join(cells) + " |")
        return lines
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        for para in shape.text_frame.paragraphs:
            text = "".join(_run_text(run) for run in para.runs).strip()
            if text:
                indent = "  " * (para.level or 0)
                lines.append(f"{indent}- {text}")
    return lines


def _title_text(slide) -> str:
    t = slide.shapes.title
    if t is not None and t.has_text_frame:
        return t.text_frame.text.strip()
    return ""


def extract_pptx(path: Path) -> str:
    prs = Presentation(str(path))
    out: list[str] = []
    for i, slide in enumerate(prs.slides, start=1):
        title = _title_text(slide)
        out.append(f"## Slide {i}" + (f": {title}" if title else ""))
        title_id = slide.shapes.title.shape_id if slide.shapes.title is not None else None
        for shape in slide.shapes:
            if title_id is not None and shape.shape_id == title_id:
                continue
            out.extend(_shape_text(shape))
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                out.append(f"Speaker notes: {notes}")
        out.append("")
    return "\n".join(out).strip()


# ---------------------------------------------------------------- diagram detection

@dataclass
class SlideInfo:
    no: int
    title: str
    hidden: bool = False
    pictures: int = 0
    charts: int = 0
    smartart: int = 0
    groups: int = 0
    connectors: int = 0
    autoshapes: int = 0

    @property
    def score(self) -> int:
        return (self.pictures * 3 + self.charts * 3 + self.smartart * 3
                + self.groups * 2 + min(self.connectors, 6) + min(self.autoshapes, 8))

    @property
    def candidate(self) -> bool:
        """Does this slide probably contain a diagram, chart or picture worth showing?"""
        return (self.pictures > 0 or self.charts > 0 or self.smartart > 0 or self.groups > 0
                or self.connectors >= 2 or self.autoshapes >= 4)


def _walk(shapes) -> Iterator:
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _walk(shape.shapes)


def _is_smartart(shape) -> bool:
    try:
        uris = shape._element.xpath(".//a:graphicData/@uri")
        return any(_DIAGRAM_URI in u for u in uris)
    except Exception:
        return False


def _picture_image(shape):
    """Return python-pptx Image for picture shapes (incl. picture placeholders) or None."""
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE or shape.__class__.__name__ in ("Picture", "PlaceholderPicture"):
        try:
            return shape.image
        except Exception:
            return None
    return None


def slide_infos(path: Path) -> list[SlideInfo]:
    prs = Presentation(str(path))
    infos = []
    for i, slide in enumerate(prs.slides, start=1):
        info = SlideInfo(no=i, title=_title_text(slide), hidden=slide._element.get("show") == "0")
        for shape in _walk(slide.shapes):
            if _picture_image(shape) is not None:
                info.pictures += 1
            elif getattr(shape, "has_chart", False) and shape.has_chart:
                info.charts += 1
            elif _is_smartart(shape):
                info.smartart += 1
            elif shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                info.groups += 1
            elif shape.__class__.__name__ == "Connector":
                info.connectors += 1
            elif shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and not shape.is_placeholder:
                info.autoshapes += 1
        infos.append(info)
    return infos


def embedded_pictures(path: Path) -> Iterator[tuple[int, int, bytes, str, tuple[int, int]]]:
    """Yield (slide_no, index, blob, ext, (w_px, h_px)) for every embedded picture."""
    prs = Presentation(str(path))
    for i, slide in enumerate(prs.slides, start=1):
        n = 0
        for shape in _walk(slide.shapes):
            img = _picture_image(shape)
            if img is None:
                continue
            n += 1
            yield i, n, img.blob, img.ext, tuple(img.size)
