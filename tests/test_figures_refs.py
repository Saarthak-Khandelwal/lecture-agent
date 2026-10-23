import io
from pathlib import Path
from types import SimpleNamespace

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches

from lecture_agent import figures as figmod
from lecture_agent.config import load_config
from lecture_agent.figures import Figure, build_figures
from lecture_agent.lectures import build_lecture
from lecture_agent.references import filter_to_search_results, urls_from_material
from lecture_agent.extractors import Source
from lecture_agent.render import render_body, used_figure_ids
from lecture_agent.writers import obsidian_writer
from lecture_agent.writers.md_to_notion import to_blocks


def make_deck(path: Path):
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[0]).shapes.title.text = "Title only"  # not a candidate
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Pipeline diagram"
    boxes = []
    for k in range(4):  # a little flowchart made of native shapes
        b = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5 + 2.2 * k), Inches(3), Inches(1.8), Inches(1))
        b.text_frame.text = f"Step {k + 1}"
        boxes.append(b)
    for a, b in zip(boxes, boxes[1:]):
        c = s.shapes.add_connector(1, a.left + a.width, a.top + a.height // 2, b.left, b.top + b.height // 2)
    # slide with a real picture
    img = Image.new("RGB", (400, 300), "steelblue")
    buf = io.BytesIO(); img.save(buf, "PNG"); buf.seek(0)
    p = prs.slides.add_slide(prs.slide_layouts[5])
    p.shapes.title.text = "A picture"
    p.shapes.add_picture(buf, Inches(1), Inches(2), Inches(4), Inches(3))
    prs.save(path)


def setup(tmp_path, monkeypatch):
    monkeypatch.delenv("OBSIDIAN_VAULT_PATH", raising=False)
    cfg = load_config()
    cfg.inbox = tmp_path / "inbox"
    cfg.figures_dir = tmp_path / "figs"
    lec = cfg.inbox / "U1" / "w1"
    lec.mkdir(parents=True)
    make_deck(lec / "deck.pptx")
    return cfg, build_lecture(lec, cfg.inbox)


def test_libreoffice_render_picks_diagram_slides(tmp_path, monkeypatch):
    cfg, lecture = setup(tmp_path, monkeypatch)
    figs = build_figures(cfg, lecture)
    ids = sorted(f.id for f in figs)
    assert ids == ["deck-s02", "deck-s03"]          # title slide skipped
    assert all(f.path.exists() and f.path.stat().st_size > 1000 for f in figs)
    assert next(f for f in figs if f.id == "deck-s02").hint == "Pipeline diagram"


def test_embedded_picture_fallback(tmp_path, monkeypatch):
    cfg, lecture = setup(tmp_path, monkeypatch)
    cfg.fig_libreoffice = False
    figs = build_figures(cfg, lecture)
    assert [f.id for f in figs] == ["deck-s03-img1"]  # shapes need rendering, pictures can be extracted


def test_pdf_pages(tmp_path):
    import pymupdf
    pdf = tmp_path / "s.pdf"
    doc = pymupdf.open()
    doc.new_page().insert_text((72, 72), "text only page")
    pg = doc.new_page()
    pg.insert_text((72, 72), "Diagram page")
    for k in range(15):
        pg.draw_rect(pymupdf.Rect(50 + k * 10, 150, 60 + k * 10, 200))
    doc.save(pdf)
    cfg = load_config()
    cfg.figures_dir = tmp_path / "f"
    lec = SimpleNamespace(unit="U", name="L", files=[pdf])
    figs = build_figures(cfg, lec)
    assert [f.id for f in figs] == ["s-p02"] and figs[0].hint == "Diagram page"


def test_obsidian_embeds_and_small_references(tmp_path, monkeypatch):
    cfg, lecture = setup(tmp_path, monkeypatch)
    figs = build_figures(cfg, lecture)
    vault = tmp_path / "vault"; vault.mkdir()
    cfg.vault_path = vault
    notes = {
        "title": "T", "summary": "s", "tags": [],
        "body_markdown": "## A\ntext\n\n{{figure:deck-s02 | Four step pipeline}}\n\n{{figure:bogus | nope}}",
        "review_questions": [],
        "references": [{"topic": "Pipelines", "title": "Intro [guide]", "url": "https://example.com/x", "note": "good overview"}],
    }
    assert used_figure_ids(notes) == ["deck-s02", "bogus"]
    out = obsidian_writer.write_note(cfg, lecture, notes, None, figs).read_text()
    assert "![[Uni/U1/attachments/w1/deck-s02.png|700]]" in out
    assert "<small>Figure: Four step pipeline (Slide 2)</small>" in out
    assert "bogus" not in out
    assert "- <small>**Pipelines**: [Intro (guide)](https://example.com/x), good overview</small>" in out
    assert (vault / "Uni/U1/attachments/w1/deck-s02.png").exists()


def test_notion_figures_and_grey_small_text(tmp_path):
    fig = Figure("a-s01", tmp_path / "a.png", "Slide 1")
    seen = []

    def resolver(fid, cap):
        seen.append((fid, cap))
        return {"object": "block", "type": "image", "image": {"caption": []}} if fid == fig.id else None

    blocks = to_blocks("intro\n\n{{figure:a-s01 | A cap}}\n\n{{figure:zzz | skipped}}\n\n- <small>**T**: [Doc](https://e.com/d)</small>", resolver)
    assert [b["type"] for b in blocks] == ["paragraph", "image", "bulleted_list_item"]
    assert seen == [("a-s01", "A cap"), ("zzz", "skipped")]
    rt = blocks[2]["bulleted_list_item"]["rich_text"]
    assert all(r["annotations"]["color"] == "gray" for r in rt)
    assert any(r["text"].get("link") for r in rt)


def test_reference_urls_are_verified_against_search_results():
    refs = [
        {"topic": "a", "title": "Real", "url": "https://docs.python.org/3/"},
        {"topic": "a", "title": "Made up", "url": "https://totally-invented.example/page"},
        {"topic": "b", "title": "Dup", "url": "https://docs.python.org/3"},
    ]
    allowed = {"https://docs.python.org/3"}
    out = filter_to_search_results(refs, allowed, 6)
    assert [r["title"] for r in out] == ["Real"]


def test_urls_from_material():
    srcs = [Source("slides", "s.pptx", "Read https://a.org/x. and (https://b.org/y)"),
            Source("transcript", "t.vtt", "https://ignored.org")]
    urls = [r["url"] for r in urls_from_material(srcs)]
    assert urls == ["https://a.org/x", "https://b.org/y"]
