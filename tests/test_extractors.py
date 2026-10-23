from pptx import Presentation

from lecture_agent.extractors import classify, extract
from lecture_agent.extractors.transcript_extractor import clean_transcript
from lecture_agent.lectures import build_lecture, lecture_dir_for

VTT = """WEBVTT

1
00:00:01.000 --> 00:00:04.000
Hello everyone

2
00:00:04.000 --> 00:00:06.000
Hello everyone

3
00:00:06.000 --> 00:00:09.000
Today we cover <c>graphs</c>
"""


def test_clean_vtt():
    assert clean_transcript(VTT, ".vtt") == "Hello everyone Today we cover graphs"


def test_classify(tmp_path):
    assert classify(tmp_path / "slides.pptx") == "slides"
    assert classify(tmp_path / "week3_transcript.txt") == "transcript"
    assert classify(tmp_path / "my notes.md") == "notes"
    assert classify(tmp_path / "board.jpg") == "image"


def test_pptx_extraction_and_grouping(tmp_path):
    inbox = tmp_path / "inbox"
    lec = inbox / "COMP1" / "week01"
    lec.mkdir(parents=True)

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Intro to Graphs"
    s.placeholders[1].text = "Nodes and edges"
    s.notes_slide.notes_text_frame.text = "Mention BFS"
    prs.save(lec / "slides.pptx")
    (lec / "notes.md").write_text("my notes")
    (lec / "~$slides.pptx").write_text("lock file")

    src = extract(lec / "slides.pptx")
    assert "Slide 1: Intro to Graphs" in src.text
    assert "Nodes and edges" in src.text
    assert "Speaker notes: Mention BFS" in src.text

    lecture = build_lecture(lec, inbox)
    assert lecture.key == "COMP1/week01"
    assert [f.name for f in lecture.files] == ["notes.md", "slides.pptx"]
    assert lecture_dir_for(lec / "slides.pptx", inbox) == lec.resolve()
    assert lecture_dir_for(inbox / "stray.pdf", inbox) is None
