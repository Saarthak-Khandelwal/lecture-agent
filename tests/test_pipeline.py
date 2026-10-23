from pathlib import Path

from pptx import Presentation

from lecture_agent import pipeline
from lecture_agent.config import load_config
from lecture_agent.state import State

FAKE = {
    "title": "Graph Basics",
    "summary": "Intro to graphs.",
    "tags": ["graphs", "bfs"],
    "body_markdown": "## Key terms\n- **Node**: a vertex\n\n## Traversal\n1. BFS\n2. DFS",
    "review_questions": ["What is a node?"],
}


def test_end_to_end_obsidian_only(tmp_path, monkeypatch):
    vault = tmp_path / "vault"
    vault.mkdir()
    (vault / "Existing.md").write_text("x")
    monkeypatch.setenv("OBSIDIAN_VAULT_PATH", str(vault))
    cfg = load_config()
    cfg.inbox = tmp_path / "inbox"
    cfg.notion_enabled = False
    cfg.references_enabled = False
    cfg.figures_dir = tmp_path / "figs"

    lec = cfg.inbox / "COMP1" / "week01"
    lec.mkdir(parents=True)
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[0]).shapes.title.text = "Graphs"
    prs.save(lec / "slides.pptx")

    calls = []

    def fake_generate(cfg_, unit, name, sources, titles, figures=None):
        calls.append(titles)
        return dict(FAKE)

    monkeypatch.setattr(pipeline, "generate_notes", fake_generate)
    state = State(tmp_path / "state.json")

    assert pipeline.process_lecture(cfg, state, lec) is True
    assert calls[0] == ["Existing"]
    out = Path(state.get("COMP1/week01")["obsidian_path"])
    text = out.read_text()
    assert out.parent == vault / "Uni" / "COMP1"
    assert text.startswith("---\n") and "# Graph Basics" in text and "## Review questions" in text

    # unchanged input is skipped, no second LLM call
    assert pipeline.process_lecture(cfg, state, lec) is False
    assert len(calls) == 1
